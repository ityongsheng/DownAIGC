import os
import re
import io
import docx
import PyPDF2
from typing import Optional

from app.doc_analyzer import StructuredParagraph, infer_heading_level


def parse_txt(file_bytes: bytes) -> str:
    # Try utf-8 first, fallback to gbk (common in Chinese environments)
    try:
        return file_bytes.decode('utf-8')
    except UnicodeDecodeError:
        try:
            return file_bytes.decode('gbk')
        except UnicodeDecodeError:
            return file_bytes.decode('utf-8', errors='ignore')


def parse_docx(file_bytes: bytes) -> str:
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
        return "\n".join([paragraph.text for paragraph in doc.paragraphs])
    except Exception as e:
        raise ValueError(f"无法解析此 Word 文档。请确保它是由正版 Office 创建的真实 .docx 文件，而不是手动修改后缀的 .doc 文件或已损坏。内部报错: {str(e)}")


def parse_docx_structured(file_bytes: bytes) -> list[StructuredParagraph]:
    """
    解析 docx 文件，返回结构化段落列表。
    每个段落包含文本、样式名、标题层级等信息。
    """
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
    except Exception as e:
        raise ValueError(
            f"无法解析此 Word 文档。请确保它是由正版 Office 创建的真实 .docx 文件，"
            f"而不是手动修改后缀的 .doc 文件或已损坏。内部报错: {str(e)}"
        )

    paragraphs = []
    for para in doc.paragraphs:
        style_name = para.style.name if para.style and para.style.name else ""
        heading_level = infer_heading_level(style_name, para.text)

        # 判断是否为列表项
        is_list_item = False
        if para.style and para.style.name:
            style_lower = para.style.name.lower()
            is_list_item = "list" in style_lower or "bullet" in style_lower

        paragraphs.append(StructuredParagraph(
            text=para.text,
            style_name=style_name,
            heading_level=heading_level,
            is_list_item=is_list_item,
        ))

    return paragraphs


def parse_text_to_structured(text: str) -> list[StructuredParagraph]:
    """
    将纯文本（来自 txt/pdf）按行转换为结构化段落列表。
    """
    lines = text.split('\n')
    paragraphs = []
    for line in lines:
        heading_level = infer_heading_level("", line)
        paragraphs.append(StructuredParagraph(
            text=line,
            style_name="",
            heading_level=heading_level,
            is_list_item=False,
        ))
    return paragraphs


def parse_pdf(file_bytes: bytes) -> str:
    try:
        reader = PyPDF2.PdfReader(io.BytesIO(file_bytes))
        text = ""
        for page in reader.pages:
            text += page.extract_text() + "\n"
        return text
    except Exception as e:
        return f"PDF解析错误: {str(e)}"
